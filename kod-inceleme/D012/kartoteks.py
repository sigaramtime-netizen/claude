# -*- coding: utf-8 -*-
"""Kartoteks (1.17) — Faz 5: salt-okunur kronolojik kartoteks (ledger) modülü.

- Cari kartoteks: seçilen cari için `cari_hareket` dökümü + koşu bakiyesi (borc − alacak).
- Stok kartoteks: seçilen stok için `stok_hareket` dökümü + koşu miktar bakiyesi (depo bazlı).
- K4: ayrı veri girişi YOK; tek kaynak `stok_hareket` / `cari_hareket`. Tüm roller salt-okunur.
- Dışa aktarma (1.17): Excel (.xlsx) / CSV + yazdırılabilir (PDF) rapor.
- Barkod hızlı sorgulama (1.17): `stok.barkod_bul()` ile doğrudan ilgili ürünün kartoteksine atlar.
"""

import datetime
import io
import urllib.parse

import db
import stok
from core import route, render_template, redirect, flash, Response

ISLEM_BADGE = {
    "Stok Girişi (Alış)": "b-ok",
    "İrsaliye Girişi": "b-ok",
    "Satış Çıkışı": "b-danger",
    "İrsaliye Çıkışı": "b-danger",
    "Servis Tüketimi": "b-warn",
    "Depolar Arası Transfer": "b-info",
}

BELGE_BADGE = {
    "Satış Faturası": "b-ok",
    "Alış Faturası": "b-danger",
    "Tahsilat": "b-info",
    "Ödeme": "b-warn",
    "Açılış Bakiyesi": "b-muted",
}

KAYNAK_LINK = {
    "Fatura": "/fatura/{}",
    "CekSenet": "/cek_senet/{}",
    "Servis": "/servis/{}",
    "Irsaliye": "/irsaliye/{}",
    "Transfer": "/stok/transferler/{}",
}
_KAYNAK_LINK_CI = {k.lower(): v for k, v in KAYNAK_LINK.items()}


def _i(v):
    try:
        return int(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _kaynak_bilgi(ilgili_modul, ilgili_kayit_id):
    """Kaynağa güvenli link (varsa) + etiket. Dönüş: (etiket, url|None).
    Modül adı büyük/küçük harfe duyarsız eşleştirilir (servis/Servis gibi)."""
    etiket = ilgili_modul or ""
    key = etiket.strip().lower()
    if key in _KAYNAK_LINK_CI and ilgili_kayit_id:
        return f"{ilgili_modul} #{ilgili_kayit_id}", _KAYNAK_LINK_CI[key].format(ilgili_kayit_id)
    return etiket, None


# --------------------------------------------------------------------------
# Ortak satır üreticiler (ekran + export + rapor aynı kaynaktan beslenir)
# --------------------------------------------------------------------------
def _hareket_detay(conn, ilgili_modul, ilgili_kayit_id):
    """Fatura/İrsaliye hareketinin kalem dökümü (detaylı görünüm için). Tek kaynak: belge."""
    modul = (ilgili_modul or "").strip().lower()
    if not ilgili_kayit_id or modul not in ("fatura", "irsaliye"):
        return []
    tablo = "fatura" if modul == "fatura" else "irsaliye"
    rows = conn.execute(
        f"SELECT k.miktar, k.birim_fiyat, k.iskonto_orani, k.kdv_orani, k.tutar, "
        f"COALESCE(k.goruntu_adi, st.ad, k.aciklama) AS gorunen_ad, "
        f"COALESCE(k.birim, st.birim, 'Adet') AS birim, st.kod AS stok_kod "
        f"FROM {tablo}_kalem k LEFT JOIN stok_kart st ON st.id=k.stok_id "
        f"WHERE k.{tablo}_id=? ORDER BY k.id", (ilgili_kayit_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def _cari_rows(conn, cari_id, bas, bit, belge_tipi, sid=None, pb=None):
    """Cari kartoteks satırları + koşu bakiye + devir. Dönüş: (rows, opening)."""
    where, params = ["cari_id=?", "sirket_id=?"], [cari_id, sid if sid is not None else 1]
    if bas:
        where.append("tarih >= ?")
        params.append(bas)
    if bit:
        where.append("tarih <= ?")
        params.append(bit)
    if belge_tipi:
        where.append("belge_tipi=?")
        params.append(belge_tipi)
    if pb:
        where.append("para_birimi=?")
        params.append(pb)
    w = " AND ".join(where)
    rows = [dict(r) for r in conn.execute(
        f"SELECT * FROM cari_hareket WHERE {w} ORDER BY tarih, id", params).fetchall()]
    opening = 0.0
    if bas:
        _ow, _op = ["cari_id=?", "tarih < ?", "sirket_id=?"], [cari_id, bas, sid if sid is not None else 1]
        if pb:
            _ow.append("para_birimi=?")
            _op.append(pb)
        opening = conn.execute(
            "SELECT COALESCE(SUM(borc-alacak),0) s FROM cari_hareket WHERE " + " AND ".join(_ow),
            _op).fetchone()["s"]
    bal = opening
    for h in rows:
        bal += (h["borc"] or 0) - (h["alacak"] or 0)
        h["bakiye"] = round(bal, 2)
        h["bakiye_cls"] = "ok" if h["bakiye"] > 0 else ("danger" if h["bakiye"] < 0 else "muted")
        h["kaynak_etiket"], h["kaynak_url"] = _kaynak_bilgi(h["ilgili_modul"], h["ilgili_kayit_id"])
    return rows, opening


def _stok_rows(conn, stok_id, depo_id, bas, bit, islem_tipi, sid=None):
    """Stok kartoteks satırları + koşu miktar + devir. Dönüş: (rows, opening)."""
    where, params = ["h.stok_id=?", "h.sirket_id=?"], [stok_id, sid if sid is not None else 1]
    if depo_id:
        where.append("h.depo_id=?")
        params.append(depo_id)
    if bas:
        where.append("h.tarih >= ?")
        params.append(bas)
    if bit:
        where.append("h.tarih <= ?")
        params.append(bit)
    if islem_tipi:
        where.append("h.islem_tipi=?")
        params.append(islem_tipi)
    w = " AND ".join(where)
    rows = [dict(r) for r in conn.execute(
        f"SELECT h.*, d.ad AS depo_ad, d.kod AS depo_kod, v.ad AS varyant_ad "
        f"FROM stok_hareket h JOIN depo d ON d.id=h.depo_id "
        f"LEFT JOIN stok_varyant v ON v.id=h.varyant_id "
        f"WHERE {w} ORDER BY h.tarih, h.id", params).fetchall()]
    opening = 0.0
    if bas:
        if depo_id:
            opening = conn.execute(
                "SELECT COALESCE(SUM(miktar),0) s FROM stok_hareket WHERE stok_id=? AND tarih < ? "
                "AND depo_id=? AND sirket_id=?", (stok_id, bas, depo_id, sid if sid is not None else 1)).fetchone()["s"]
        else:
            opening = conn.execute(
                "SELECT COALESCE(SUM(miktar),0) s FROM stok_hareket WHERE stok_id=? AND tarih < ? "
                "AND sirket_id=?", (stok_id, bas, sid if sid is not None else 1)).fetchone()["s"]
    kalan = opening
    for h in rows:
        m = h["miktar"] or 0
        h["giris"] = round(max(m, 0), 2)
        h["cikis"] = round(max(-m, 0), 2)
        kalan += m
        h["kalan"] = round(kalan, 2)
        h["kalan_cls"] = "ok" if h["kalan"] > 0 else ("danger" if h["kalan"] < 0 else "muted")
        h["kaynak_etiket"], h["kaynak_url"] = _kaynak_bilgi(h["ilgili_modul"], h["ilgili_kayit_id"])
        c = stok.hareket_cari(conn, h["ilgili_modul"], h["ilgili_kayit_id"], sid)
        h["cari_unvan"] = c["cari_unvan"] if c else None
        h["cari_kod"] = c["cari_kod"] if c else None
        h["cari_id"] = c["cari_id"] if c else None
    return rows, opening


# --------------------------------------------------------------------------
# Dışa aktarma yardımcıları (Stok2 deseni — K4 ile aynı kaynaktan)
# --------------------------------------------------------------------------
def _export_csv(basliklar, veri, dosya_adi):
    buf = io.StringIO()
    buf.write("\ufeff")  # UTF-8 BOM — Excel Türkçe karakterleri doğru açsın
    buf.write(";".join(basliklar) + "\r\n")
    for satir in veri:
        buf.write(";".join(str(x) if x is not None else "" for x in satir) + "\r\n")
    return Response(buf.getvalue(), "200 OK", [
        ("Content-Type", "text/csv; charset=utf-8"),
        ("Content-Disposition", f'attachment; filename="{dosya_adi}"'),
    ])


def _export_xlsx(basliklar, veri, dosya_adi):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Kartoteks"
        ws.append(basliklar)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        for satir in veri:
            ws.append(satir)
        from openpyxl.utils import get_column_letter
        for i in range(1, len(basliklar) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 22 if basliklar[i - 1] in ("Açıklama", "Kaynak") else 14
        buf = io.BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(), "200 OK", [
            ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("Content-Disposition", f'attachment; filename="{dosya_adi}"'),
        ])
    except Exception:  # noqa: BLE001 — openpyxl yoksa None (çağıran CSV'ye düşer)
        return None


# --------------------------------------------------------------------------
# Rotalar
# --------------------------------------------------------------------------
@route(r"/kartoteks", roles=())
def kartoteks_index(req):
    conn = db.get_conn()
    sid = db.sirket_id(req)
    n_cari = conn.execute("SELECT COUNT(*) c FROM cari_kart WHERE aktif=1 AND sirket_id=?",
                          (sid,)).fetchone()["c"]
    n_stok = conn.execute("SELECT COUNT(*) c FROM stok_kart WHERE aktif=1 AND sirket_id=?",
                          (sid,)).fetchone()["c"]
    n_cari_h = conn.execute("SELECT COUNT(*) c FROM cari_hareket WHERE sirket_id=?", (sid,)).fetchone()["c"]
    n_stok_h = conn.execute("SELECT COUNT(*) c FROM stok_hareket WHERE sirket_id=?", (sid,)).fetchone()["c"]
    toplam_stok = conn.execute("SELECT COALESCE(SUM(miktar),0) s FROM stok_seviye WHERE sirket_id=?",
                               (sid,)).fetchone()["s"]
    conn.close()
    return render_template("kartoteks/index.html", n_cari=n_cari, n_stok=n_stok,
                           n_cari_h=n_cari_h, n_stok_h=n_stok_h, toplam_stok=toplam_stok)


@route(r"/kartoteks/cari", roles=())
def kartoteks_cari(req):
    conn = db.get_conn()
    q = req.q("q").strip()
    cari_id = _i(req.q("cari_id"))
    bas = req.q("bas")
    bit = req.q("bit")
    belge_tipi = req.q("belge_tipi")
    gorunum = req.q("gorunum") if req.q("gorunum") in ("detay", "ozet") else "ozet"
    fiyat_gizli = req.q("fiyat") == "gizle"
    kdv_haric = req.q("kdv") == "haric"
    pb = req.q("pb")
    pbs = db.para_birimleri(sid=db.sirket_id(req))
    if pb and pb not in pbs:
        pb = None

    cariler = conn.execute("SELECT id, kod, unvan, tip FROM cari_kart WHERE aktif=1 AND sirket_id=? ORDER BY kod",
                            (db.sirket_id(req),)).fetchall()
    if q:
        cariler = [c for c in cariler if q.lower() in (c["kod"] + c["unvan"]).lower()]
    tipler = sorted({r["belge_tipi"] for r in conn.execute(
        "SELECT DISTINCT belge_tipi FROM cari_hareket WHERE sirket_id=?", (db.sirket_id(req),)).fetchall()})

    cari = None
    rows = []
    opening = 0.0
    top_b = top_a = 0.0
    bakiye = 0.0
    filt_top_b = filt_top_a = 0.0
    pb_totals = []
    if cari_id:
        c = conn.execute("SELECT * FROM cari_kart WHERE id=? AND sirket_id=?",
                         (cari_id, db.sirket_id(req))).fetchone()
        if c:
            cari = dict(c)
            rows, opening = _cari_rows(conn, cari_id, bas, bit, belge_tipi, db.sirket_id(req), pb)
            # KDV dahil/hariç + detaylı görünüm zenginleştirmesi (yalnız görünüm; cari_hareket değişmez).
            fat_ids = {h["ilgili_kayit_id"] for h in rows
                       if (h["ilgili_modul"] or "").strip().lower() == "fatura" and h["ilgili_kayit_id"]}
            kdv_map = {}
            if fat_ids:
                _ph = ",".join("?" * len(fat_ids))
                kdv_map = {r["id"]: (r["kdv_toplam"] or 0.0) for r in conn.execute(
                    f"SELECT id, kdv_toplam FROM fatura WHERE id IN ({_ph})", tuple(fat_ids)).fetchall()}
            for h in rows:
                modul = (h["ilgili_modul"] or "").strip().lower()
                kdv_t = kdv_map.get(h["ilgili_kayit_id"], 0.0) if modul == "fatura" else 0.0
                h["kdv_tutar"] = kdv_t
                b = h["borc"] or 0.0
                a = h["alacak"] or 0.0
                if kdv_haric and modul == "fatura":
                    if b:
                        b -= kdv_t
                    if a:
                        a -= kdv_t
                h["g_borc"] = round(b, 2)
                h["g_alacak"] = round(a, 2)
                h["detay"] = _hareket_detay(conn, h["ilgili_modul"], h["ilgili_kayit_id"]) if gorunum == "detay" else []
            bal = opening
            for h in rows:
                bal += h["g_borc"] - h["g_alacak"]
                h["bakiye"] = round(bal, 2)
                h["bakiye_cls"] = "ok" if h["bakiye"] > 0 else ("danger" if h["bakiye"] < 0 else "muted")
            filt_top_b = round(sum(h["g_borc"] for h in rows), 2)
            filt_top_a = round(sum(h["g_alacak"] for h in rows), 2)
            tot = conn.execute(
                "SELECT COALESCE(SUM(borc),0) b, COALESCE(SUM(alacak),0) a FROM cari_hareket "
                "WHERE cari_id=? AND sirket_id=?", (cari_id, db.sirket_id(req))).fetchone()
            top_b, top_a = tot["b"], tot["a"]
            bakiye = round(top_b - top_a, 2)
            # D012-D — döviz bazlı alt toplamlar (filtreden bağımsız özet kartlar).
            pb_totals = []
            for _pb in pbs:
                pr = conn.execute(
                    "SELECT COALESCE(SUM(borc),0) b, COALESCE(SUM(alacak),0) a, "
                    "AVG(CASE WHEN doviz_kur>0 THEN doviz_kur END) k FROM cari_hareket "
                    "WHERE cari_id=? AND sirket_id=? AND para_birimi=?",
                    (cari_id, db.sirket_id(req), _pb)).fetchone()
                pb_totals.append({"kod": _pb, "borc": pr["b"], "alacak": pr["a"],
                                  "net": (pr["b"] or 0) - (pr["a"] or 0), "kur": pr["k"] or 0.0})
        else:
            pb_totals = []
    conn.close()
    return render_template(
        "kartoteks/cari.html", cariler=cariler, cari=cari, rows=rows, opening=opening,
        top_b=top_b, top_a=top_a, bakiye=bakiye,
        filt_top_b=filt_top_b, filt_top_a=filt_top_a,
        pbs=pbs, pb=pb, pb_totals=pb_totals,
        tipler=tipler, BELGE_BADGE=BELGE_BADGE, gorunum=gorunum,
        fiyat_gizli=fiyat_gizli, kdv_haric=kdv_haric,
        filtro={"q": q, "cari_id": cari_id, "bas": bas, "bit": bit, "belge_tipi": belge_tipi,
                "gorunum": gorunum, "fiyat": "gizle" if fiyat_gizli else "goster",
                "kdv": "haric" if kdv_haric else "dahil", "pb": pb})


@route(r"/kartoteks/stok", roles=())
def kartoteks_stok(req):
    conn = db.get_conn()

    # Barkod hızlı sorgulama (1.17): okutulan barkod doğrudan ürünün kartoteksine atlar.
    barkod = req.q("barkod").strip()
    if barkod:
        import stok
        sonuc = stok.barkod_bul(conn, barkod, db.sirket_id(req))
        if sonuc:
            conn.close()
            flash(req, f"Barkod {barkod} → {sonuc['kart']['ad']}", "info")
            return redirect(f"/kartoteks/stok?stok_id={sonuc['stok_id']}")
        flash(req, f"Barkod bulunamadı: {barkod}", "danger")

    q = req.q("q").strip()
    stok_id = _i(req.q("stok_id"))
    depo_id = _i(req.q("depo_id"))
    bas = req.q("bas")
    bit = req.q("bit")
    islem_tipi = req.q("islem_tipi")

    stoklar = conn.execute("SELECT id, kod, ad, barkod, birim FROM stok_kart WHERE aktif=1 AND sirket_id=? ORDER BY kod",
                            (db.sirket_id(req),)).fetchall()
    if q:
        stoklar = [s for s in stoklar if q.lower() in (s["kod"] + s["ad"] + (s["barkod"] or "")).lower()]
    tipler = sorted({r["islem_tipi"] for r in conn.execute(
        "SELECT DISTINCT islem_tipi FROM stok_hareket WHERE sirket_id=?", (db.sirket_id(req),)).fetchall()})
    depolar = conn.execute("SELECT id, ad FROM depo WHERE aktif=1 AND sirket_id=? ORDER BY id",
                           (db.sirket_id(req),)).fetchall()

    kart = None
    rows = []
    opening = 0.0
    toplam_mevcut = 0.0
    dagilim = []
    if stok_id:
        k = conn.execute("SELECT * FROM stok_kart WHERE id=? AND sirket_id=?",
                         (stok_id, db.sirket_id(req))).fetchone()
        if k:
            kart = dict(k)
            dagilim = [dict(r) for r in conn.execute(
                "SELECT d.id, d.ad, COALESCE(sv.miktar,0) miktar FROM depo d "
                "LEFT JOIN stok_seviye sv ON sv.depo_id=d.id AND sv.stok_id=? AND sv.sirket_id=? "
                "WHERE d.aktif=1 AND d.sirket_id=? ORDER BY d.id",
                (stok_id, db.sirket_id(req), db.sirket_id(req))).fetchall()]
            toplam_mevcut = round(sum(d["miktar"] for d in dagilim), 2)
            rows, opening = _stok_rows(conn, stok_id, depo_id, bas, bit, islem_tipi, db.sirket_id(req))
    conn.close()
    return render_template(
        "kartoteks/stok.html", stoklar=stoklar, kart=kart, rows=rows, opening=opening,
        toplam_mevcut=toplam_mevcut, dagilim=dagilim, tipler=tipler, depolar=depolar,
        ISLEM_BADGE=ISLEM_BADGE,
        filtro={"q": q, "barkod": barkod, "stok_id": stok_id, "depo_id": depo_id,
                "bas": bas, "bit": bit, "islem_tipi": islem_tipi})


# --------------------------------------------------------------------------
# Dışa aktarma (Excel / CSV)
# --------------------------------------------------------------------------
@route(r"/kartoteks/cari/export", roles=())
def kartoteks_cari_export(req):
    conn = db.get_conn()
    cari_id = _i(req.q("cari_id"))
    cari = conn.execute("SELECT * FROM cari_kart WHERE id=? AND sirket_id=?",
                        (cari_id or 0, db.sirket_id(req))).fetchone() if cari_id else None
    if not cari:
        conn.close()
        return redirect("/kartoteks/cari")
    rows, opening = _cari_rows(conn, cari_id, req.q("bas"), req.q("bit"), req.q("belge_tipi"),
                               db.sirket_id(req))
    conn.close()

    fmt = req.q("format") or "xlsx"
    basliklar = ["Tarih", "Belge Tipi", "Belge No", "Açıklama", "Kaynak", "Vade", "Borç", "Alacak", "Bakiye"]
    veri = [[h["tarih"], h["belge_tipi"], h["belge_no"] or "", h["aciklama"] or "",
             h["kaynak_etiket"] or "", h["vade"] or "", h["borc"] or 0, h["alacak"] or 0, h["bakiye"]]
            for h in rows]
    if opening:
        veri.insert(0, ["", "Önceki dönemden devir", "", "", "", "", "", "", opening])
    if fmt == "csv":
        return _export_csv(basliklar, veri, f"kartoteks_cari_{cari['kod']}.csv")
    r = _export_xlsx(basliklar, veri, f"kartoteks_cari_{cari['kod']}.xlsx")
    return r if r is not None else _export_csv(basliklar, veri, f"kartoteks_cari_{cari['kod']}.csv")


@route(r"/kartoteks/stok/export", roles=())
def kartoteks_stok_export(req):
    conn = db.get_conn()
    stok_id = _i(req.q("stok_id"))
    kart = conn.execute("SELECT * FROM stok_kart WHERE id=? AND sirket_id=?",
                        (stok_id or 0, db.sirket_id(req))).fetchone() if stok_id else None
    if not kart:
        conn.close()
        return redirect("/kartoteks/stok")
    rows, opening = _stok_rows(conn, stok_id, _i(req.q("depo_id")), req.q("bas"), req.q("bit"),
                               req.q("islem_tipi"), db.sirket_id(req))
    conn.close()

    fmt = req.q("format") or "xlsx"
    basliklar = ["Tarih", "İşlem Tipi", "Depo", "Belge No", "Açıklama", "Kaynak", "Giriş", "Çıkış", "Kalan"]
    veri = [[h["tarih"], h["islem_tipi"], h["depo_ad"], h["belge_no"] or "", h["aciklama"] or "",
             h["kaynak_etiket"] or "", h["giris"], h["cikis"], h["kalan"]] for h in rows]
    if opening:
        veri.insert(0, ["", "Önceki dönemden devir", "", "", "", "", "", "", opening])
    if fmt == "csv":
        return _export_csv(basliklar, veri, f"kartoteks_stok_{kart['kod']}.csv")
    r = _export_xlsx(basliklar, veri, f"kartoteks_stok_{kart['kod']}.xlsx")
    return r if r is not None else _export_csv(basliklar, veri, f"kartoteks_stok_{kart['kod']}.csv")


# --------------------------------------------------------------------------
# Yazdırılabilir rapor (PDF'e kaydet — Stok2 "rapor" deseni)
# --------------------------------------------------------------------------
@route(r"/kartoteks/cari/rapor", roles=())
def kartoteks_cari_rapor(req):
    conn = db.get_conn()
    cari_id = _i(req.q("cari_id"))
    cari = conn.execute("SELECT * FROM cari_kart WHERE id=? AND sirket_id=?",
                        (cari_id or 0, db.sirket_id(req))).fetchone() if cari_id else None
    if not cari:
        conn.close()
        return redirect("/kartoteks/cari")
    rows, opening = _cari_rows(conn, cari_id, req.q("bas"), req.q("bit"), req.q("belge_tipi"),
                               db.sirket_id(req))
    conn.close()
    return render_template("kartoteks/cari_rapor.html", cari=cari, rows=rows, opening=opening,
                           bugun=datetime.date.today().isoformat(),
                           filtro={"bas": req.q("bas"), "bit": req.q("bit"), "belge_tipi": req.q("belge_tipi")})


@route(r"/kartoteks/stok/rapor", roles=())
def kartoteks_stok_rapor(req):
    conn = db.get_conn()
    stok_id = _i(req.q("stok_id"))
    kart = conn.execute("SELECT * FROM stok_kart WHERE id=? AND sirket_id=?",
                        (stok_id or 0, db.sirket_id(req))).fetchone() if stok_id else None
    if not kart:
        conn.close()
        return redirect("/kartoteks/stok")
    rows, opening = _stok_rows(conn, stok_id, _i(req.q("depo_id")), req.q("bas"), req.q("bit"),
                               req.q("islem_tipi"), db.sirket_id(req))
    conn.close()
    return render_template("kartoteks/stok_rapor.html", kart=kart, rows=rows, opening=opening,
                           bugun=datetime.date.today().isoformat(),
                           filtro={"depo_id": req.q("depo_id"), "bas": req.q("bas"),
                                   "bit": req.q("bit"), "islem_tipi": req.q("islem_tipi")})


def register():
    return "kartoteks"
